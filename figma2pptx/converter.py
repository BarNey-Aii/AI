"""Figma node tree -> editable PPTX."""

from __future__ import annotations

import io
import math
from dataclasses import dataclass
from xml.sax.saxutils import quoteattr

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.util import Emu

from . import geometry as g
from .figma_api import AssetProvider
from .paint import EMU_PER_PX, effects_xml, fill_xml, line_xml, visible_paints
from .text import text_body_xml

NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

CONTAINER_TYPES = {"FRAME", "COMPONENT", "COMPONENT_SET", "INSTANCE", "SECTION", "SLIDE", "GROUP",
                   "SLIDE_ROW", "SLIDE_GRID", "INTERACTIVE_SLIDE_ELEMENT"}
FRAME_LIKE = CONTAINER_TYPES - {"GROUP"}
VECTOR_TYPES = {"VECTOR", "STAR", "REGULAR_POLYGON", "POLYGON", "BOOLEAN_OPERATION", "ELLIPSE",
                "WASHI_TAPE", "HIGHLIGHT"}
KNOWN_TYPES = CONTAINER_TYPES | VECTOR_TYPES | {"RECTANGLE", "TEXT", "LINE"}

SAFE_IMAGE_FORMATS = {"PNG", "JPEG", "GIF", "BMP", "TIFF"}


@dataclass
class Options:
    slide_width_in: float | None = None   # None = keep the Figma frame size 1:1 (96 dpi)
    rasterize_gradients: bool = False     # True = gradients become exact PNG picture fills
    font_weights: str = "bold"            # "bold" or "names" (e.g. "Inter SemiBold")
    raster_fallback: bool = True          # masks/unknown nodes rendered via Figma API as PNG
    lang: str = "cs-CZ"


class _Ctx:
    def __init__(self, emu: float, rasterize_gradients: bool):
        self.emu = emu
        self.rasterize_gradients = rasterize_gradients
        self.slide = None
        self._img_cache: dict[int, str] = {}

    def embed(self, data: bytes) -> str:
        key = hash(data)
        if key in self._img_cache:
            return self._img_cache[key]
        data = _normalize_image(data)
        _part, r_id = self.slide.part.get_or_add_image_part(io.BytesIO(data))
        self._img_cache[key] = r_id
        return r_id


def _normalize_image(data: bytes) -> bytes:
    try:
        with Image.open(io.BytesIO(data)) as im:
            if im.format in SAFE_IMAGE_FORMATS:
                return data
            buf = io.BytesIO()
            im.save(buf, "PNG")
            return buf.getvalue()
    except Exception:
        return data


class _El:
    __slots__ = ("xml", "bounds")

    def __init__(self, xml: str, bounds):
        self.xml = xml
        self.bounds = bounds  # (x, y, w, h) in px


def _node_size(node: dict) -> tuple[float, float]:
    size = node.get("size")
    if size:
        return float(size.get("x", 0)), float(size.get("y", 0))
    bb = node.get("absoluteBoundingBox") or {}
    return float(bb.get("width", 0)), float(bb.get("height", 0))


def _radii(node: dict) -> list[float]:
    r = node.get("rectangleCornerRadii")
    if r and len(r) == 4:
        return [float(v) for v in r]
    c = float(node.get("cornerRadius") or 0)
    return [c, c, c, c]


class Converter:
    def __init__(self, assets: AssetProvider | None = None, options: Options | None = None):
        self.assets = assets or AssetProvider()
        self.opt = options or Options()
        self.warnings: list[str] = []
        self._rendered: dict[str, bytes] = {}

    # ------------------------------------------------------------------ public
    def convert(self, slide_nodes: list[dict]) -> Presentation:
        if not slide_nodes:
            raise ValueError("Nebyly nalezeny žádné snímky / rámce k převodu.")
        fw, fh = _node_size(slide_nodes[0])
        fw, fh = fw or 1920, fh or 1080
        scale = 1.0
        if self.opt.slide_width_in:
            scale = self.opt.slide_width_in * 914400 / (fw * EMU_PER_PX)
        # PowerPoint limit: 56 in per side.
        max_side = 51206400
        if max(fw, fh) * EMU_PER_PX * scale > max_side:
            scale = max_side / (max(fw, fh) * EMU_PER_PX)
        self.ctx = _Ctx(EMU_PER_PX * scale, self.opt.rasterize_gradients)
        self.lang = self.opt.lang
        self.font_weights = self.opt.font_weights

        prs = Presentation()
        prs.slide_width = Emu(int(round(fw * self.ctx.emu)))
        prs.slide_height = Emu(int(round(fh * self.ctx.emu)))
        layout = prs.slide_layouts[6]

        if self.opt.raster_fallback and self.assets.can_render():
            ids: list[str] = []
            for s in slide_nodes:
                self._collect_raster(s, ids, is_root=True)
            if ids:
                self._rendered = self.assets.render_nodes(ids)

        for node in slide_nodes:
            slide = prs.slides.add_slide(layout)
            self._convert_slide(slide, node)
        return prs

    # ------------------------------------------------------------------ helpers
    def _needs_raster(self, node: dict) -> bool:
        t = node.get("type")
        if t not in KNOWN_TYPES:
            return True
        if t in CONTAINER_TYPES or t == "BOOLEAN_OPERATION":
            return any(c.get("isMask") and c.get("visible", True) for c in node.get("children") or [])
        return False

    def _collect_raster(self, node: dict, ids: list[str], is_root=False) -> None:
        if not node.get("visible", True):
            return
        if not is_root and self._needs_raster(node):
            ids.append(node["id"])
            return
        if node.get("type") == "BOOLEAN_OPERATION":
            return
        for c in node.get("children") or []:
            self._collect_raster(c, ids)

    def hyperlink(self, url: str) -> str:
        return self.ctx.slide.part.relate_to(url, RT.HYPERLINK, is_external=True)

    def _next_id(self) -> int:
        self._id += 1
        return self._id

    def _xfrm(self, pl: g.Placement, tag: str = "a:xfrm") -> str:
        e = self.ctx.emu
        attrs = ""
        if pl.rot:
            attrs += f' rot="{int(round(pl.rot * 60000)) % 21600000}"'
        if pl.flip_h:
            attrs += ' flipH="1"'
        return (f'<{tag}{attrs}><a:off x="{int(round(pl.x * e))}" y="{int(round(pl.y * e))}"/>'
                f'<a:ext cx="{max(0, int(round(pl.w * e)))}" cy="{max(0, int(round(pl.h * e)))}"/></{tag}>')

    def _sp(self, name: str, pl: g.Placement, geom: str, fill: str | None, ln: str | None = None,
            effects: str = "", txbody: str | None = None, txbox: bool = False) -> _El:
        sid = self._next_id()
        cnv = '<p:cNvSpPr txBox="1"/>' if txbox else "<p:cNvSpPr/>"
        fill = fill or "<a:noFill/>"
        body = txbody or ""
        xml = (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name={quoteattr(name)}/>{cnv}<p:nvPr/></p:nvSpPr>'
               f"<p:spPr>{self._xfrm(pl)}{geom}{fill}{ln or '<a:ln><a:noFill/></a:ln>'}{effects}</p:spPr>"
               f"{body}</p:sp>")
        return _El(xml, pl.bounds())

    def _pic(self, name: str, pl: g.Placement, r_id: str) -> _El:
        sid = self._next_id()
        xml = (f'<p:pic><p:nvPicPr><p:cNvPr id="{sid}" name={quoteattr(name)}/>'
               f'<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr/></p:nvPicPr>'
               f'<p:blipFill><a:blip r:embed="{r_id}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
               f'<p:spPr>{self._xfrm(pl)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
        return _El(xml, pl.bounds())

    def _grp(self, name: str, els: list[_El]) -> _El | None:
        els = [e for e in els if e]
        if not els:
            return None
        if len(els) == 1:
            return els[0]
        x0 = min(e.bounds[0] for e in els)
        y0 = min(e.bounds[1] for e in els)
        x1 = max(e.bounds[0] + e.bounds[2] for e in els)
        y1 = max(e.bounds[1] + e.bounds[3] for e in els)
        e = self.ctx.emu
        off = f'x="{int(round(x0 * e))}" y="{int(round(y0 * e))}"'
        ext = f'cx="{int(round((x1 - x0) * e))}" cy="{int(round((y1 - y0) * e))}"'
        sid = self._next_id()
        xml = (f'<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{sid}" name={quoteattr(name)}/><p:cNvGrpSpPr/><p:nvPr/>'
               f'</p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off {off}/><a:ext {ext}/><a:chOff {off}/>'
               f'<a:chExt {ext}/></a:xfrm></p:grpSpPr>{"".join(x.xml for x in els)}</p:grpSp>')
        return _El(xml, (x0, y0, x1 - x0, y1 - y0))

    def _image_bytes(self, paint: dict) -> bytes | None:
        if paint.get("type") != "IMAGE":
            return None
        ref = paint.get("imageRef") or paint.get("gifRef")
        if not ref:
            return None
        data = self.assets.image_fill(ref)
        if data is None:
            self.warnings.append(f"Obrázek {ref} se nepodařilo stáhnout.")
        return data

    # ------------------------------------------------------------------ transforms
    def _transform(self, node: dict, parent_m, container_m) -> g.Matrix:
        rel = node.get("relativeTransform")
        w, h = _node_size(node)
        bb = node.get("absoluteBoundingBox")
        target = None
        if bb:
            target = (float(bb["x"]) - self._origin[0], float(bb["y"]) - self._origin[1],
                      float(bb["width"]), float(bb["height"]))
        if not rel:
            if target:
                return g.translate(target[0], target[1])
            return parent_m
        rel = g.mat(rel)
        cands = [g.mul(container_m, rel)]
        if parent_m is not container_m:
            cands.append(g.mul(parent_m, rel))
        if target is None or len(cands) == 1:
            return cands[0]

        def err(m):
            b = g.aabb(m, 0, 0, w, h)
            return sum(abs(a - c) for a, c in zip(b, target))

        return min(cands, key=err)

    # ------------------------------------------------------------------ slides
    def _convert_slide(self, slide, root: dict) -> None:
        self.ctx.slide = slide
        self.ctx._img_cache = {}
        self._id = 1
        bb = root.get("absoluteBoundingBox") or {"x": 0, "y": 0}
        self._origin = (float(bb.get("x", 0)), float(bb.get("y", 0)))
        w, h = _node_size(root)
        root_m = g.IDENTITY
        rel = root.get("relativeTransform")
        if rel and root.get("type") not in FRAME_LIKE:
            m = g.mat(rel)
            lin = ((m[0][0], m[0][1], 0.0), (m[1][0], m[1][1], 0.0))
            bx, by, _, _ = g.aabb(lin, 0, 0, w, h)
            root_m = ((m[0][0], m[0][1], -bx), (m[1][0], m[1][1], -by))
        els: list[_El] = []
        opacity = float(root.get("opacity", 1.0))

        if root.get("type") in FRAME_LIKE:
            fills = visible_paints(root.get("fills"))
            if fills and not any(_radii(root)):
                bg = fill_xml(fills[0], self.ctx, w, h, (0, 0, w, h), opacity, self._image_bytes(fills[0]))
                if bg:
                    self._set_background(slide, bg)
                    fills = fills[1:]
            for i, paint in enumerate(fills):
                fx = fill_xml(paint, self.ctx, w, h, (0, 0, w, h), opacity, self._image_bytes(paint))
                if fx:
                    els.append(self._sp(f"{root.get('name', 'Pozadí')} – pozadí {i + 1}",
                                        g.Placement(0, 0, w, h), g.prst_geom_xml("rect"), fx))
            container_m = root_m
        else:
            container_m = g.translate(-self._origin[0], -self._origin[1])

        for child in root.get("children") or []:
            el = self._node(child, root_m, container_m, opacity)
            if el:
                els.append(el)
        self._append(slide, els)

    def _set_background(self, slide, fill: str) -> None:
        bg = parse_xml(f"<p:bg {NS}><p:bgPr>{fill}<a:effectLst/></p:bgPr></p:bg>")
        c_sld = slide._element.cSld
        old = c_sld.find("{http://schemas.openxmlformats.org/presentationml/2006/main}bg")
        if old is not None:
            c_sld.remove(old)
        c_sld.insert(0, bg)

    def _append(self, slide, els: list[_El]) -> None:
        if not els:
            return
        wrapper = parse_xml(f"<p:spTree {NS}>{''.join(e.xml for e in els)}</p:spTree>")
        sp_tree = slide.shapes._spTree
        ext = sp_tree.find("{http://schemas.openxmlformats.org/presentationml/2006/main}extLst")
        for child in list(wrapper):
            if ext is not None:
                ext.addprevious(child)
            else:
                sp_tree.append(child)

    # ------------------------------------------------------------------ nodes
    def _node(self, node: dict, parent_m, container_m, opacity: float) -> _El | None:
        if not node.get("visible", True) or node.get("isMask"):
            return None
        t = node.get("type")
        m = self._transform(node, parent_m, container_m)
        op = opacity * float(node.get("opacity", 1.0))
        name = node.get("name") or t or "Tvar"

        if node.get("id") in self._rendered:
            return self._raster(node, name)
        if self._needs_raster(node):
            if t not in KNOWN_TYPES:
                self.warnings.append(f"Nepodporovaný typ uzlu {t} ({name}) byl vynechán.")
                return None
            self.warnings.append(f"Maska v '{name}' není v PPTX podporovaná – vykresleno bez masky.")

        if t in CONTAINER_TYPES:
            els: list[_El] = []
            if t in FRAME_LIKE:
                els.extend(self._box_layers(node, m, op, name, "rect"))
            child_container = m if t != "GROUP" else container_m
            for c in node.get("children") or []:
                el = self._node(c, m, child_container, op)
                if el:
                    els.append(el)
            return self._grp(name, els)
        if t == "RECTANGLE":
            return self._grp(name, self._box_layers(node, m, op, name, "rect"))
        if t == "ELLIPSE" and _full_ellipse(node):
            return self._grp(name, self._box_layers(node, m, op, name, "ellipse"))
        if t == "TEXT":
            return self._text(node, m, op, name)
        if t == "LINE":
            return self._line(node, m, op, name)
        if t in VECTOR_TYPES:
            return self._grp(name, self._vector_layers(node, m, op, name))
        return None

    def _raster(self, node: dict, name: str) -> _El | None:
        data = self._rendered.get(node["id"])
        bb = node.get("absoluteBoundingBox")
        if not data or not bb:
            return None
        pl = g.Placement(float(bb["x"]) - self._origin[0], float(bb["y"]) - self._origin[1],
                         float(bb["width"]), float(bb["height"]))
        return self._pic(f"{name} (obrázek)", pl, self.ctx.embed(data))

    def _box_layers(self, node: dict, m, op: float, name: str, kind: str) -> list[_El]:
        w, h = _node_size(node)
        fills = visible_paints(node.get("fills"))
        strokes = visible_paints(node.get("strokes"))
        sw = float(node.get("strokeWeight") or 0)
        if not strokes or sw <= 0:
            strokes = []
        effects = effects_xml(node, self.ctx, op)
        if not fills and not strokes and not effects:
            return []
        if strokes and node.get("individualStrokeWeights"):
            iw = node["individualStrokeWeights"]
            self.warnings.append(f"'{name}': rozdílné tloušťky okrajů nejsou v PPTX podporované.")
            sw = max(float(iw.get(k, 0)) for k in ("top", "right", "bottom", "left"))
        radii = _radii(node)

        def geom(d: float) -> str:
            if kind == "ellipse":
                return g.prst_geom_xml("ellipse")
            r = [max(0.0, v + d) if v > 0 else 0.0 for v in radii]
            bw, bh = w + 2 * d, h + 2 * d
            if not any(r):
                return g.prst_geom_xml("rect")
            if max(r) - min(r) < 0.01:
                adj = int(round(min(50000, r[0] / max(1e-6, min(bw, bh)) * 100000)))
                return g.prst_geom_xml("roundRect", {"adj": adj})
            return g.cust_geom_xml([g.rounded_rect_path(bw, bh, r)], 0, 0, bw, bh, self.ctx.emu)

        align = node.get("strokeAlign", "CENTER")
        d = {"INSIDE": -sw / 2, "OUTSIDE": sw / 2}.get(align, 0.0) if strokes else 0.0
        stroke_paint = strokes[-1] if strokes else None
        opaque_stroke = bool(stroke_paint) and stroke_paint.get("type") == "SOLID" and \
            float(stroke_paint.get("color", {}).get("a", 1)) * float(stroke_paint.get("opacity", 1)) * op >= 0.999
        separate_stroke = bool(strokes) and (d != 0 and not opaque_stroke)

        layers: list[_El] = []
        fill_xmls = []
        for p in fills:
            fx = fill_xml(p, self.ctx, w, h, (0, 0, w, h), op, self._image_bytes(p))
            if fx:
                fill_xmls.append((p, fx))
        if not fill_xmls:
            fill_xmls = [(None, "<a:noFill/>")] if strokes or effects else []
        for i, (p, fx) in enumerate(fill_xmls):
            top = i == len(fill_xmls) - 1
            lname = name if i == 0 else f"{name} – výplň {i + 1}"
            eff = effects if i == 0 else ""
            if top and strokes and not separate_stroke:
                box = (-d, -d, w + 2 * d, h + 2 * d)
                if p is not None and d:
                    fx = fill_xml(p, self.ctx, w, h, box, op, self._image_bytes(p)) or fx
                ln = line_xml(node, self.ctx, w, h, box, op, weight=sw)
                pl = g.place(m, *box)
                layers.append(self._sp(lname, pl, geom(d), fx, ln, eff))
            else:
                layers.append(self._sp(lname, g.place(m, 0, 0, w, h), geom(0), fx, None, eff))
        if separate_stroke:
            box = (-d, -d, w + 2 * d, h + 2 * d)
            ln = line_xml(node, self.ctx, w, h, box, op, weight=sw)
            layers.append(self._sp(f"{name} – okraj", g.place(m, *box), geom(d), None, ln))
        return layers

    def _vector_layers(self, node: dict, m, op: float, name: str) -> list[_El]:
        w, h = _node_size(node)
        fills = visible_paints(node.get("fills"))
        strokes = visible_paints(node.get("strokes"))
        sw = float(node.get("strokeWeight") or 0)
        effects = effects_xml(node, self.ctx, op)
        fill_paths = [g.parse_svg_path(p.get("path", "")) for p in node.get("fillGeometry") or []]
        stroke_paths = [g.parse_svg_path(p.get("path", "")) for p in node.get("strokeGeometry") or []]
        fill_paths = [p for p in fill_paths if p]
        stroke_paths = [p for p in stroke_paths if p]
        layers: list[_El] = []

        def box_for(paths):
            bb = g.segments_bbox(paths)
            x0, y0, x1, y1 = 0.0, 0.0, w, h
            if bb:
                x0, y0 = min(x0, bb[0]), min(y0, bb[1])
                x1, y1 = max(x1, bb[2]), max(y1, bb[3])
            return x0, y0, max(x1 - x0, 1e-3), max(y1 - y0, 1e-3)

        if fill_paths and fills:
            box = box_for(fill_paths)
            geom = g.cust_geom_xml(fill_paths, box[0], box[1], box[2], box[3], self.ctx.emu)
            fxs = [fx for fx in (fill_xml(p, self.ctx, w, h, box, op, self._image_bytes(p)) for p in fills) if fx]
            for i, fx in enumerate(fxs):
                ln = None
                if i == len(fxs) - 1 and strokes and sw > 0 and not stroke_paths:
                    ln = line_xml(node, self.ctx, w, h, box, op, weight=sw)
                layers.append(self._sp(name if i == 0 else f"{name} – výplň {i + 1}", g.place(m, *box),
                                       geom, fx, ln, effects if i == 0 else ""))
        if strokes and sw > 0:
            if stroke_paths:
                box = box_for(stroke_paths)
                geom = g.cust_geom_xml(stroke_paths, box[0], box[1], box[2], box[3], self.ctx.emu)
                for i, p in enumerate(strokes):
                    fx = fill_xml(p, self.ctx, w, h, box, op, self._image_bytes(p))
                    if fx:
                        layers.append(self._sp(f"{name} – okraj" if i == 0 else f"{name} – okraj {i + 1}",
                                               g.place(m, *box), geom, fx, None,
                                               effects if not layers else ""))
            elif not (fill_paths and fills) and fill_paths:
                box = box_for(fill_paths)
                geom = g.cust_geom_xml(fill_paths, box[0], box[1], box[2], box[3], self.ctx.emu, fill=False)
                layers.append(self._sp(name, g.place(m, *box), geom, None,
                                       line_xml(node, self.ctx, w, h, box, op, weight=sw), effects))
        if not layers and not fill_paths and not stroke_paths and (fills or strokes):
            self.warnings.append(f"'{name}': chybí geometrie vektoru (stáhněte s geometry=paths).")
        return layers

    def _line(self, node: dict, m, op: float, name: str) -> _El | None:
        w, _h = _node_size(node)
        strokes = visible_paints(node.get("strokes"))
        if not strokes:
            return None
        pl = g.place(m, 0, 0, w, 0)
        ln = line_xml(node, self.ctx, w, 1, (0, 0, w, 1), op)
        return self._sp(name, pl, g.prst_geom_xml("line"), None, ln, effects_xml(node, self.ctx, op))

    def _text(self, node: dict, m, op: float, name: str) -> _El | None:
        if not node.get("characters"):
            return None
        w, h = _node_size(node)
        pl = g.place(m, 0, 0, w, h)
        body = text_body_xml(node, self, op, w, h)
        return self._sp(name, pl, g.prst_geom_xml("rect"), None, None, "", body, txbox=True)


def _full_ellipse(node: dict) -> bool:
    arc = node.get("arcData")
    if not arc:
        return True
    start = float(arc.get("startingAngle", 0))
    end = float(arc.get("endingAngle", 2 * math.pi))
    inner = float(arc.get("innerRadius", 0))
    return inner == 0 and abs((end - start) - 2 * math.pi) < 1e-3


def to_xml_string(el) -> str:  # debugging helper
    return etree.tostring(el, pretty_print=True).decode()
