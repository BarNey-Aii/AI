import io
import os
import unittest
import zipfile

from pptx import Presentation

from figma2pptx import Options, convert_json, parse_figma_url
from figma2pptx.geometry import parse_svg_path, place, IDENTITY
from figma2pptx.paint import linear_gradient_xml

HERE = os.path.dirname(__file__)
FIXTURE = os.path.join(HERE, "fixtures", "sample_slide.json")
IMAGES = os.path.join(HERE, "fixtures", "images")


class UrlTests(unittest.TestCase):
    def test_design_link(self):
        link = parse_figma_url("https://www.figma.com/design/AbC123/My-Deck?node-id=12-34&t=x")
        self.assertEqual(link.file_key, "AbC123")
        self.assertEqual(link.node_ids, ["12:34"])

    def test_slides_and_branch_link(self):
        self.assertEqual(parse_figma_url("https://www.figma.com/slides/KEY9/Deck?node-id=1%3A2").node_ids, ["1:2"])
        self.assertEqual(parse_figma_url("https://figma.com/design/MAIN/branch/BR1/X").file_key, "BR1")

    def test_invalid(self):
        with self.assertRaises(Exception):
            parse_figma_url("https://example.com/nothing")


class GeometryTests(unittest.TestCase):
    def test_svg_path(self):
        segs = parse_svg_path("M0 0L10 0C10 5 5 10 0 10Q-5 5 0 0Z")
        self.assertEqual([s[0] for s in segs], ["M", "L", "C", "Q", "Z"])

    def test_relative_path(self):
        segs = parse_svg_path("m1 1 l2 0 h3 v4 z")
        self.assertEqual(segs[1], ("L", (3.0, 1.0)))
        self.assertEqual(segs[3], ("L", (6.0, 5.0)))

    def test_place_rotation(self):
        import math
        r = math.radians(30)
        m = ((math.cos(r), -math.sin(r), 0.0), (math.sin(r), math.cos(r), 0.0))
        pl = place(m, 0, 0, 200, 100)
        self.assertAlmostEqual(pl.rot, 30.0, places=4)
        self.assertAlmostEqual(pl.w, 200)
        self.assertEqual(place(IDENTITY, 5, 6, 7, 8).x, 5)


class GradientTests(unittest.TestCase):
    def test_horizontal_linear(self):
        paint = {"type": "GRADIENT_LINEAR",
                 "gradientHandlePositions": [{"x": 0, "y": 0.5}, {"x": 1, "y": 0.5}, {"x": 0, "y": 1}],
                 "gradientStops": [{"position": 0, "color": {"r": 1, "g": 0, "b": 0, "a": 1}},
                                   {"position": 1, "color": {"r": 0, "g": 0, "b": 1, "a": 1}}]}
        xml = linear_gradient_xml(paint, 400, 100, (0, 0, 400, 100), 1.0)
        self.assertIn('ang="0"', xml)
        self.assertIn('val="FF0000"', xml)
        self.assertIn('val="0000FF"', xml)

    def test_partial_linear_is_remapped(self):
        paint = {"type": "GRADIENT_LINEAR",
                 "gradientHandlePositions": [{"x": 0.25, "y": 0.5}, {"x": 0.75, "y": 0.5}, {"x": 0.25, "y": 1}],
                 "gradientStops": [{"position": 0, "color": {"r": 0, "g": 0, "b": 0, "a": 1}},
                                   {"position": 1, "color": {"r": 1, "g": 1, "b": 1, "a": 1}}]}
        xml = linear_gradient_xml(paint, 400, 100, (0, 0, 400, 100), 1.0)
        self.assertIn('pos="25000"', xml)
        self.assertIn('pos="75000"', xml)


class ConvertTests(unittest.TestCase):
    def test_fixture_converts(self):
        data, warnings = convert_json(FIXTURE, IMAGES, Options())
        prs = Presentation(io.BytesIO(data))
        self.assertEqual(prs.slide_width, 1920 * 9525)
        names = [s.name for s in prs.slides[0].shapes]
        self.assertIn("Karta", names)
        self.assertIn("Skupina", names)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("ppt/slides/slide1.xml").decode()
            self.assertTrue(any(n.startswith("ppt/media/") for n in z.namelist()))
        self.assertIn("<p:bg>", xml)
        self.assertIn("a:gradFill", xml)
        self.assertIn("a:custGeom", xml)
        self.assertIn("a:outerShdw", xml)
        self.assertIn("světe", xml)

    def test_slide_width_option_and_raster_gradients(self):
        data, _ = convert_json(FIXTURE, IMAGES, Options(slide_width_in=10, rasterize_gradients=True))
        prs = Presentation(io.BytesIO(data))
        self.assertEqual(prs.slide_width, 9144000)
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("ppt/slides/slide1.xml").decode()
        # Shape gradients become picture fills; text gradients stay native.
        self.assertIn("a:blipFill", xml)


if __name__ == "__main__":
    unittest.main()
